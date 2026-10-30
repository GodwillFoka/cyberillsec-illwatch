"""Inventaire en fichier : import et export Excel ou CSV (M8, ADR-028)."""

import io
import zipfile
from pathlib import Path
from uuid import uuid4

import pytest
from click.testing import CliRunner
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from illwatch.app.security import create_access_token
from illwatch.modules.correlation.assets import AssetInput, create_asset, find_asset
from illwatch.modules.correlation.inventory_file import (
    MAX_FILE_BYTES,
    InventoryFileError,
    import_inventory,
    inventory_csv,
    inventory_xlsx,
    read_inventory,
    template_xlsx,
)
from illwatch.modules.foundation.users import create_user
from illwatch.shared.enums import Severity, UserRole

HEADER = "nom;produit;editeur;version;plateforme;criticite;adresses;domaines;notes;en_service"


def _xlsx(rows: list[list[object]], sheet: str = "Inventaire") -> bytes:
    workbook = Workbook()
    ws = workbook.active
    assert ws is not None
    ws.title = sheet
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _headers(session: AsyncSession, role: UserRole) -> dict[str, str]:
    name = f"{role.value.lower()}-{uuid4().hex[:8]}"
    user = await create_user(
        session,
        username=name,
        email=f"{name}@cyberill.test",
        password="mot-de-passe-robuste-2026",
        role=role,
    )
    return {"Authorization": f"Bearer {create_access_token(user.id, user.role)}"}


# --- Lecture -------------------------------------------------------------------------------


def test_csv_point_virgule_libelles_francais() -> None:
    content = (
        HEADER + "\n"
        "Pare-feu siège;FortiOS;Fortinet;7.2.5;;Critique;198.51.100.10, 198.51.100.11;"
        "vpn.example.com;Cluster;oui\n"
        "Wiki;Confluence;Atlassian;8.5;Debian 12;élevée;;wiki.example.com;;non\n"
        ";;;;;;;;;\n"
    ).encode()
    parsed = read_inventory(content, "inventaire.csv")
    assert not parsed.errors
    first, second = (row.data for row in parsed.rows)
    assert first is not None and second is not None
    assert first.criticality is Severity.CRITICAL
    assert first.addresses == ("198.51.100.10", "198.51.100.11")
    assert first.is_active is True
    assert second.criticality is Severity.HIGH and second.platform == "Debian 12"
    assert second.is_active is False
    assert len(parsed.rows) == 2  # ligne vide ignorée


def test_csv_virgule_windows_1252_et_entetes_souples() -> None:
    content = 'Name,Product,Vendor,Criticité\n"Messagerie siège",Exchange Server,Microsoft,HIGH\n'
    parsed = read_inventory(content.encode("cp1252"), "export.csv")
    assert not parsed.errors
    row = parsed.rows[0]
    assert row.name == "Messagerie siège"
    assert row.data is not None and row.data.vendor == "Microsoft"
    assert parsed.columns == frozenset({"nom", "produit", "editeur", "criticite"})


def test_xlsx_nombres_et_onglet_inventaire() -> None:
    content = _xlsx([["nom", "produit", "version"], ["VPN", "FortiOS", 7.0], ["Proxy", "Squid", 6]])
    parsed = read_inventory(content, "inventaire.xlsx")
    versions = [row.data.version for row in parsed.rows if row.data]
    assert versions == ["7", "6"]


def test_lignes_refusees_avec_motif() -> None:
    content = (
        b"nom;produit;criticite;adresses;domaines\n"
        b"A;OS;Moyenne;;\n"  # produit trop court
        b"B;FortiOS;Urgente;;\n"  # criticité inconnue
        b"C;FortiOS;;999.1.1.1;\n"  # adresse invalide
        b"D;FortiOS;;;pas un domaine !\n"
        b"E;FortiOS;;;\n"
        b"e;FortiOS;;;\n"  # doublon (casse ignorée)
    )
    parsed = read_inventory(content)
    errors = {row.line: row.error for row in parsed.errors}
    assert set(errors) == {2, 3, 4, 5, 7}
    assert "Produit obligatoire" in (errors[2] or "")
    assert "Critique, Élevée, Moyenne ou Faible" in (errors[3] or "")
    assert "ligne 6" in (errors[7] or "")


@pytest.mark.parametrize(
    ("content", "message"),
    [
        (b"", "vide"),
        (b"produit;version\nFortiOS;7\n", "Manquant : nom"),
        (b"PK\x03\x04pas une archive", "illisible"),
        (b"x" * (MAX_FILE_BYTES + 1), "trop volumineux"),
    ],
    # Identifiants explicites : sans eux, pytest nomme le test d'après le contenu (2 Mo de
    # « x » dans un seul nom), ce qui bloque l'affichage du journal de la CI.
    ids=["vide", "colonne-manquante", "archive-illisible", "trop-volumineux"],
)
def test_fichiers_refuses_en_bloc(content: bytes, message: str) -> None:
    with pytest.raises(InventoryFileError, match=message):
        read_inventory(content, "inventaire.xlsx" if content.startswith(b"PK") else "a.csv")


def test_archive_excel_trop_lourde_refusee_avant_lecture() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (45 * 1024 * 1024))
    with pytest.raises(InventoryFileError, match="décompressé"):
        read_inventory(buffer.getvalue(), "bombe.xlsx")


# --- Import en base --------------------------------------------------------------------------


async def test_apercu_puis_import_et_colonnes_absentes_conservees(db_session: AsyncSession) -> None:
    await create_asset(
        db_session,
        AssetInput(
            name="Wiki",
            product="Confluence",
            notes="Note à garder",
            addresses=("198.51.100.20",),
            criticality=Severity.LOW,
        ),
    )
    parsed = read_inventory(b"nom;produit;version\nWiki;Confluence;8.5.3\nVPN;FortiOS;7.2\n")

    preview = await import_inventory(db_session, parsed, apply=False)
    assert not preview.applied
    assert [r.action for r in preview.rows] == ["update", "create"]
    assert await find_asset(db_session, "VPN") is None

    report = await import_inventory(db_session, parsed, apply=True)
    assert report.applied and (report.created, report.updated) == (1, 1)
    wiki = await find_asset(db_session, "Wiki")
    assert wiki is not None
    assert wiki.version == "8.5.3"
    assert wiki.notes == "Note à garder"  # colonne absente : valeur conservée
    assert wiki.addresses == ["198.51.100.20"]
    assert wiki.criticality == "LOW"


async def test_tout_ou_rien(db_session: AsyncSession) -> None:
    parsed = read_inventory(b"nom;produit\nBon;FortiOS\nMauvais;OS\n")
    report = await import_inventory(db_session, parsed, apply=True)
    assert not report.applied and report.rejected == 1
    assert await find_asset(db_session, "Bon") is None


# --- Export ----------------------------------------------------------------------------------


async def test_export_reimportable_et_formules_neutralisees(db_session: AsyncSession) -> None:
    asset = await create_asset(
        db_session,
        AssetInput(
            name='=HYPERLINK("http://evil")',
            product="FortiOS",
            criticality=Severity.CRITICAL,
            domains=("vpn.example.com",),
        ),
    )
    text = inventory_csv([asset]).decode("utf-8-sig")
    assert text.splitlines()[0] == HEADER
    assert "'=HYPERLINK" in text.splitlines()[1]  # apostrophe : Excel n'évalue pas la formule

    workbook = load_workbook(io.BytesIO(inventory_xlsx([asset])))
    sheet = workbook["Inventaire"]
    assert sheet["A2"].value.startswith("'=")
    assert sheet["F2"].value == "Critique"

    again = read_inventory(inventory_xlsx([asset]), "inventaire.xlsx")
    assert not again.errors and again.rows[0].data is not None
    assert again.rows[0].data.criticality is Severity.CRITICAL


def test_modele_excel() -> None:
    workbook = load_workbook(io.BytesIO(template_xlsx()))
    assert workbook.sheetnames == ["Inventaire", "Mode d'emploi"]
    assert [c.value for c in workbook["Inventaire"][1]] == HEADER.split(";")
    with pytest.raises(InventoryFileError, match="Aucune ligne|Manquant"):
        # Le modèle vide ne contient que l'en-tête : rien à importer, mais il est lisible.
        read_inventory(b"\n", "vide.csv")
    assert read_inventory(template_xlsx(), "modele.xlsx").rows == []


def test_commande_template(tmp_path: Path) -> None:
    from illwatch.cli.correlate import assets

    target = tmp_path / "modele.xlsx"
    result = CliRunner().invoke(assets, ["template", "-o", str(target)])
    assert result.exit_code == 0, result.output
    assert load_workbook(target).sheetnames[0] == "Inventaire"


# --- API -------------------------------------------------------------------------------------


async def test_api_import_export_et_droits(client: AsyncClient, db_session: AsyncSession) -> None:
    viewer = await _headers(db_session, UserRole.VIEWER)
    analyst = await _headers(db_session, UserRole.ANALYST)
    content = _xlsx([["nom", "produit", "criticite"], ["Messagerie", "Exchange Server", "Élevée"]])
    files = {"file": ("inventaire.xlsx", content, "application/octet-stream")}

    denied = await client.post("/api/v1/assets/import", files=files, headers=viewer)
    assert denied.status_code == 403

    preview = await client.post("/api/v1/assets/import", files=files, headers=analyst)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["applied"] is False and body["rows"][0]["action"] == "create"
    assert body["columns"] == ["nom", "produit", "criticite"]

    done = await client.post(
        "/api/v1/assets/import", params={"apply": "true"}, files=files, headers=analyst
    )
    assert done.status_code == 200 and done.json()["created"] == 1
    listed = (await client.get("/api/v1/assets", headers=viewer)).json()
    assert [a["name"] for a in listed] == ["Messagerie"]

    bad = {"file": ("a.csv", b"produit\nFortiOS\n", "text/csv")}
    refused = await client.post("/api/v1/assets/import", files=bad, headers=analyst)
    assert refused.status_code == 422 and "nom" in refused.json()["detail"]

    exported = await client.get("/api/v1/assets/export", params={"format": "csv"}, headers=analyst)
    assert exported.status_code == 200
    assert "attachment" in exported.headers["content-disposition"]
    assert "Messagerie;Exchange Server" in exported.content.decode("utf-8-sig")
    assert (await client.get("/api/v1/assets/export", headers=viewer)).status_code == 403

    template = await client.get("/api/v1/assets/template", headers=viewer)
    assert template.status_code == 200
    assert template.headers["content-type"].startswith("application/vnd.openxmlformats")
