"""Inventaire d'actifs en fichier : import et export CSV ou Excel (M8, ADR-028).

Un seul analyseur sert l'écran Inventaire, l'API et la commande `illwatch assets import` :

- **formats** : Excel (`.xlsx`) et CSV ; séparateur `;`, `,` ou tabulation, détecté sur la ligne
  d'en-tête ; encodage UTF-8 (avec ou sans BOM) ou Windows-1252, comme les exports d'Excel ;
- **en-têtes souples** : casse, accents et espaces ignorés, synonymes anglais acceptés
  (`name`, `product`, `vendor`…) ;
- **criticité** en clair (« Critique », « Élevée », « Moyenne », « Faible ») ou par code
  (`CRITICAL`, `HIGH`, `MEDIUM`, `LOW`) ;
- **tout ou rien** : une seule ligne invalide et rien n'est écrit ; l'aperçu liste chaque ligne
  avec l'action prévue (création, mise à jour) ou le motif du refus ;
- **colonnes absentes conservées** : un fichier sans colonne « notes » ne vide pas les notes des
  actifs existants.

Garde-fous : taille du fichier, nombre de lignes, archive Excel contrôlée avant lecture (taille
décompressée, nombre de membres) ; à l'export, les cellules qui commencent par `= + - @` sont
neutralisées (injection de formule, CWE-1236).
"""

import csv
import io
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field, replace
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from illwatch.app.models import Asset
from illwatch.modules.correlation.assets import (
    AssetError,
    AssetInput,
    create_asset,
    find_asset,
    update_asset,
    validate_asset,
)
from illwatch.shared.enums import Severity

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000
# Archive Excel : un classeur d'inventaire légitime pèse quelques centaines de kilo-octets
# décompressé ; au-delà, c'est une erreur ou une « bombe » de décompression.
MAX_XLSX_UNCOMPRESSED = 40 * 1024 * 1024
MAX_XLSX_MEMBERS = 500

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NAME = "Inventaire"

# Colonnes, dans l'ordre du modèle et des exports. Seuls « nom » et « produit » sont obligatoires.
COLUMNS: tuple[str, ...] = (
    "nom",
    "produit",
    "editeur",
    "version",
    "plateforme",
    "criticite",
    "adresses",
    "domaines",
    "notes",
    "en_service",
)
REQUIRED = ("nom", "produit")

HEADER_LABELS: dict[str, str] = {
    "nom": "nom",
    "produit": "produit",
    "editeur": "editeur",
    "version": "version",
    "plateforme": "plateforme",
    "criticite": "criticite",
    "adresses": "adresses",
    "domaines": "domaines",
    "notes": "notes",
    "en_service": "en_service",
}

# En-têtes reconnus, une fois normalisés (minuscules, sans accent ni séparateur).
_SYNONYMS: dict[str, str] = {
    "nom": "nom",
    "name": "nom",
    "actif": "nom",
    "asset": "nom",
    "nomdelactif": "nom",
    "produit": "produit",
    "product": "produit",
    "editeur": "editeur",
    "vendor": "editeur",
    "fabricant": "editeur",
    "constructeur": "editeur",
    "version": "version",
    "plateforme": "plateforme",
    "platform": "plateforme",
    "systeme": "plateforme",
    "os": "plateforme",
    "criticite": "criticite",
    "criticality": "criticite",
    "adresses": "adresses",
    "adresse": "adresses",
    "adressesip": "adresses",
    "adresseip": "adresses",
    "ip": "adresses",
    "addresses": "adresses",
    "domaines": "domaines",
    "domaine": "domaines",
    "domains": "domaines",
    "notes": "notes",
    "note": "notes",
    "commentaire": "notes",
    "commentaires": "notes",
    "enservice": "en_service",
    "active": "en_service",
    "isactive": "en_service",
    "inservice": "en_service",
}

CRITICALITY_LABELS: dict[Severity, str] = {
    Severity.CRITICAL: "Critique",
    Severity.HIGH: "Élevée",
    Severity.MEDIUM: "Moyenne",
    Severity.LOW: "Faible",
}

_CRITICALITY_WORDS: dict[str, Severity] = {
    "critical": Severity.CRITICAL,
    "critique": Severity.CRITICAL,
    "high": Severity.HIGH,
    "elevee": Severity.HIGH,
    "eleve": Severity.HIGH,
    "haute": Severity.HIGH,
    "haut": Severity.HIGH,
    "medium": Severity.MEDIUM,
    "moyenne": Severity.MEDIUM,
    "moyen": Severity.MEDIUM,
    "low": Severity.LOW,
    "faible": Severity.LOW,
    "basse": Severity.LOW,
    "bas": Severity.LOW,
}

_TRUE = {"oui", "o", "yes", "y", "vrai", "true", "1", "x", "enservice", "actif"}
_FALSE = {"non", "n", "no", "faux", "false", "0", "retire", "horsservice", "inactif"}

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


class InventoryFileError(ValueError):
    """Fichier illisible dans son ensemble (format, taille, en-tête) : message pour l'analyste."""


def _key(text: str) -> str:
    """Forme comparable : minuscules, sans accent, sans espace ni ponctuation."""
    decomposed = unicodedata.normalize("NFKD", text.strip().lower())
    return "".join(c for c in decomposed if c.isalnum())


def _cell(value: Any) -> str:
    """Valeur de cellule en texte : 7.0 (nombre Excel) devient « 7 »."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def split_values(text: str) -> tuple[str, ...]:
    """Adresses ou domaines séparés par espaces, virgules, points-virgules ou retours à la ligne."""
    for separator in (",", ";", "\n", "\r", "\t"):
        text = text.replace(separator, " ")
    return tuple(v for v in text.split() if v)


def parse_criticality(text: str) -> Severity:
    if not text.strip():
        return Severity.MEDIUM
    level = _CRITICALITY_WORDS.get(_key(text))
    if level is None:
        raise AssetError(
            f"criticité « {text.strip()} » inconnue : Critique, Élevée, Moyenne ou Faible."
        )
    return level


def parse_flag(text: str) -> bool:
    key = _key(text)
    if not key or key in _TRUE:
        return True
    if key in _FALSE:
        return False
    raise AssetError(f"« en_service » vaut « {text.strip()} » : oui ou non attendu.")


@dataclass(frozen=True)
class InventoryRow:
    """Une ligne du fichier, validée ou refusée."""

    line: int
    name: str
    data: AssetInput | None = None
    error: str | None = None
    action: str = ""  # « create » ou « update », fixé par l'aperçu


@dataclass
class InventoryFile:
    """Contenu lu : colonnes présentes et lignes (données vides ignorées)."""

    columns: frozenset[str]
    rows: list[InventoryRow] = field(default_factory=list)

    @property
    def errors(self) -> list[InventoryRow]:
        return [r for r in self.rows if r.error]


@dataclass(frozen=True)
class ImportReport:
    applied: bool
    created: int
    updated: int
    rows: list[InventoryRow]

    @property
    def rejected(self) -> int:
        return sum(1 for r in self.rows if r.error)


# --- Lecture -------------------------------------------------------------------------------


def _is_xlsx(content: bytes, filename: str) -> bool:
    return content[:4] == b"PK\x03\x04" or filename.lower().endswith((".xlsx", ".xlsm"))


def _xlsx_table(content: bytes) -> list[list[str]]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise InventoryFileError(
            "Fichier Excel illisible : enregistrez-le au format .xlsx (ou CSV) puis réessayez."
        ) from None
    members = archive.infolist()
    if len(members) > MAX_XLSX_MEMBERS or sum(m.file_size for m in members) > MAX_XLSX_UNCOMPRESSED:
        raise InventoryFileError("Classeur Excel trop volumineux une fois décompressé.")
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lève des exceptions variées sur un classeur abîmé
        raise InventoryFileError(f"Classeur Excel illisible ({type(exc).__name__}).") from None
    try:
        names = workbook.sheetnames
        sheet = workbook[SHEET_NAME] if SHEET_NAME in names else workbook.worksheets[0]
        table: list[list[str]] = []
        for values in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in values]
            if any(cells):  # lignes vides ou seulement mises en forme : ignorées
                table.append(cells)
                if len(table) > MAX_ROWS + 1:
                    break
        return table
    finally:
        workbook.close()


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def _csv_table(content: bytes) -> list[list[str]]:
    text = _decode(content)
    first = text.split("\n", 1)[0]
    delimiter = max((";", ",", "\t"), key=first.count)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    table: list[list[str]] = []
    for values in reader:
        cells = [v.strip() for v in values]
        if any(cells):
            table.append(cells)
            if len(table) > MAX_ROWS + 1:
                break
    return table


def read_inventory(content: bytes, filename: str = "") -> InventoryFile:
    """Lit et valide un fichier d'inventaire, sans toucher à la base."""
    if not content:
        raise InventoryFileError("Fichier vide.")
    if len(content) > MAX_FILE_BYTES:
        raise InventoryFileError(
            f"Fichier trop volumineux ({len(content) // 1024} Ko) : "
            f"{MAX_FILE_BYTES // (1024 * 1024)} Mo au plus."
        )
    table = _xlsx_table(content) if _is_xlsx(content, filename) else _csv_table(content)
    if not table:
        raise InventoryFileError("Aucune ligne dans le fichier.")

    header = [_SYNONYMS.get(_key(cell), "") for cell in table[0]]
    present = frozenset(h for h in header if h)
    missing = [c for c in REQUIRED if c not in present]
    if missing:
        raise InventoryFileError(
            "En-tête attendu sur la première ligne : "
            + ";".join(COLUMNS)
            + f" (au moins nom et produit). Manquant : {', '.join(missing)}."
        )
    if len(table) - 1 > MAX_ROWS:
        raise InventoryFileError(f"{MAX_ROWS} actifs au plus par fichier.")

    parsed = InventoryFile(columns=present)
    seen: dict[str, int] = {}
    for number, values in enumerate(table[1:], start=2):
        cells = {
            column: values[index] if index < len(values) else ""
            for index, column in enumerate(header)
            if column
        }
        name = " ".join(cells.get("nom", "").split())
        try:
            data = AssetInput(
                name=name,
                product=cells.get("produit", ""),
                vendor=cells.get("editeur") or None,
                version=cells.get("version") or None,
                platform=cells.get("plateforme") or None,
                criticality=parse_criticality(cells.get("criticite", "")),
                addresses=split_values(cells.get("adresses", "")),
                domains=split_values(cells.get("domaines", "")),
                notes=cells.get("notes") or None,
                is_active=parse_flag(cells.get("en_service", "")),
            )
            validate_asset(data)
            if name.casefold() in seen:
                raise AssetError(f"actif déjà présent à la ligne {seen[name.casefold()]}.")
            seen[name.casefold()] = number
        except AssetError as exc:
            message = str(exc)
            parsed.rows.append(InventoryRow(number, name, error=message[:1].upper() + message[1:]))
            continue
        parsed.rows.append(InventoryRow(number, name, data=data))
    return parsed


# --- Écriture en base ------------------------------------------------------------------------


def _merge(existing: Asset, data: AssetInput, columns: frozenset[str]) -> AssetInput:
    """Garde la valeur actuelle de chaque colonne absente du fichier."""
    keep: dict[str, Any] = {}
    if "editeur" not in columns:
        keep["vendor"] = existing.vendor
    if "version" not in columns:
        keep["version"] = existing.version
    if "plateforme" not in columns:
        keep["platform"] = existing.platform
    if "criticite" not in columns:
        keep["criticality"] = Severity(existing.criticality)
    if "adresses" not in columns:
        keep["addresses"] = tuple(existing.addresses or ())
    if "domaines" not in columns:
        keep["domains"] = tuple(existing.domains or ())
    if "notes" not in columns:
        keep["notes"] = existing.notes
    if "en_service" not in columns:
        keep["is_active"] = existing.is_active
    return replace(data, **keep)


async def import_inventory(
    session: AsyncSession, parsed: InventoryFile, *, apply: bool
) -> ImportReport:
    """Aperçu (`apply=False`) ou import (`apply=True`, seulement si aucune ligne n'est refusée).

    Un actif est reconnu par son nom : présent, il est mis à jour ; absent, il est créé.
    """
    rows: list[InventoryRow] = []
    for row in parsed.rows:
        if row.error or row.data is None:
            rows.append(row)
            continue
        existing = await find_asset(session, row.name)
        rows.append(replace(row, action="update" if existing is not None else "create"))

    if not apply or any(r.error for r in rows):
        return ImportReport(applied=False, created=0, updated=0, rows=rows)

    created = updated = 0
    for row in rows:
        assert row.data is not None
        existing = await find_asset(session, row.name)
        try:
            if existing is None:
                await create_asset(session, row.data)
                created += 1
            else:
                await update_asset(session, existing.id, _merge(existing, row.data, parsed.columns))
                updated += 1
        except AssetError as exc:  # course rare : un actif homonyme créé entre l'aperçu et l'import
            raise InventoryFileError(f"Ligne {row.line} : {exc}") from exc
    return ImportReport(applied=True, created=created, updated=updated, rows=rows)


# --- Export et modèle ------------------------------------------------------------------------


def _safe(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def asset_cells(asset: Asset) -> list[str]:
    return [
        asset.name,
        asset.product,
        asset.vendor or "",
        asset.version or "",
        asset.platform or "",
        CRITICALITY_LABELS[Severity(asset.criticality)],
        " ".join(asset.addresses or []),
        " ".join(asset.domains or []),
        asset.notes or "",
        "oui" if asset.is_active else "non",
    ]


def inventory_csv(assets: Iterable[Asset]) -> bytes:
    """CSV « ; » en UTF-8 avec BOM : Excel l'ouvre directement avec les accents."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([HEADER_LABELS[c] for c in COLUMNS])
    for asset in assets:
        writer.writerow([_safe(v) for v in asset_cells(asset)])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


GUIDE_ROWS: tuple[tuple[str, str, str, str], ...] = (
    ("nom", "Obligatoire", "Nom unique de l'actif dans votre organisation.", "Pare-feu siège"),
    (
        "produit",
        "Obligatoire",
        "Nom du produit tel qu'il apparaît dans les avis de sécurité (3 caractères au moins). "
        "Évitez les noms trop vagues comme « Windows ».",
        "FortiOS",
    ),
    ("editeur", "Conseillé", "Éditeur ou fabricant.", "Fortinet"),
    (
        "version",
        "Conseillé",
        "Version installée : sans elle, la corrélation ne peut pas écarter les versions non "
        "vulnérables. Saisie en texte (7.10 reste 7.10).",
        "7.2.5",
    ),
    (
        "plateforme",
        "Facultatif",
        "Système hôte : tranche les CVE propres à une distribution.",
        "Debian 12",
    ),
    (
        "criticite",
        "Facultatif",
        "Critique, Élevée, Moyenne ou Faible (Moyenne si vide).",
        "Critique",
    ),
    (
        "adresses",
        "Facultatif",
        "Adresses IP publiques exposées sur Internet, séparées par des espaces ou des virgules.",
        "198.51.100.10 198.51.100.11",
    ),
    (
        "domaines",
        "Facultatif",
        "Domaines exposés ; un domaine couvre ses sous-domaines.",
        "vpn.example.com",
    ),
    ("notes", "Facultatif", "Commentaire libre (2 000 caractères au plus).", "Cluster principal"),
    (
        "en_service",
        "Facultatif",
        "oui ou non (oui si vide). Un actif hors service n'est plus corrélé.",
        "oui",
    ),
)


def _xlsx(assets: Sequence[Asset], *, with_guide: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    night, signal = "0A1628", "22D3EE"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = SHEET_NAME
    sheet.append([HEADER_LABELS[c] for c in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=night)
    for asset in assets:
        sheet.append([_safe(v) for v in asset_cells(asset)])
    widths = (26, 22, 16, 12, 18, 12, 34, 30, 30, 11)
    for index, width in enumerate(widths, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = width
        # Texte pour toute la colonne : Excel ne transforme ni « 7.10 » en 7,1 ni « 0012 » en 12.
        for row in range(2, max(len(assets) + 2, 501)):
            sheet.cell(row=row, column=index).number_format = "@"
    sheet.freeze_panes = "A2"
    choices = DataValidation(
        type="list", formula1='"Critique,Élevée,Moyenne,Faible"', allow_blank=True
    )
    choices.add("F2:F5001")
    service = DataValidation(type="list", formula1='"oui,non"', allow_blank=True)
    service.add("J2:J5001")
    sheet.add_data_validation(choices)
    sheet.add_data_validation(service)

    if with_guide:
        guide = workbook.create_sheet("Mode d'emploi")
        guide.append(["ILL Watch : modèle d'inventaire des actifs"])
        guide["A1"].font = Font(bold=True, size=14, color=night)
        guide.append(
            [
                "Remplissez l'onglet « Inventaire », une ligne par actif, puis importez le "
                "fichier depuis l'écran Inventaire (bouton « Importer ») ou avec "
                "« illwatch assets import inventaire.xlsx »."
            ]
        )
        guide.append([])
        guide.append(["Colonne", "Statut", "Contenu attendu", "Exemple"])
        for cell in guide[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=night)
        for entry in GUIDE_ROWS:
            guide.append(list(entry))
        guide.append([])
        guide.append(
            [
                "Un actif déjà présent (même nom) est mis à jour ; une colonne absente du "
                "fichier ne modifie pas la valeur existante. Si une ligne est refusée, rien "
                "n'est importé : l'aperçu indique le motif de chaque refus."
            ]
        )
        for letter, width in zip("ABCD", (16, 13, 80, 28), strict=True):
            guide.column_dimensions[letter].width = width
        for row in guide.iter_rows(min_row=5):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        guide["A1"].fill = PatternFill("solid", fgColor=signal)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def inventory_xlsx(assets: Sequence[Asset]) -> bytes:
    return _xlsx(assets, with_guide=False)


def template_xlsx() -> bytes:
    """Modèle vide à remplir : en-têtes, listes de choix, onglet « Mode d'emploi »."""
    return _xlsx([], with_guide=True)


def template_csv() -> bytes:
    return inventory_csv([])
